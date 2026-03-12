#!/usr/bin/env python3
"""
Brief → Word Document Generator — Nexus BDR Agent
====================================================
Takes a brief JSON file and generates a polished Word document.
Handles incomplete briefs gracefully (skips empty phases).

Usage:
    python3 brief_to_docx.py --brief outputs/briefs/brief_mellow_fellow_*.json
    python3 brief_to_docx.py --brief outputs/briefs/brief_mellow_fellow_*.json --output custom_name.docx
    python3 brief_to_docx.py --all   # Convert all briefs in outputs/briefs/

Requires: pip install python-docx
"""

import os, sys, json, re, argparse, glob
from datetime import datetime
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Inches, Pt, Cm, RGBColor, Emu
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
except ImportError:
    print("Installing python-docx...")
    os.system(f"{sys.executable} -m pip install python-docx --quiet")
    from docx import Document
    from docx.shared import Inches, Pt, Cm, RGBColor, Emu
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn

SCRIPT_DIR = Path(__file__).parent.resolve()
OUTPUT_DIR = SCRIPT_DIR / "outputs" if (SCRIPT_DIR / "outputs").exists() else SCRIPT_DIR.parent / "outputs"
BRIEFS_DIR = OUTPUT_DIR / "briefs"

GOLD = RGBColor(0xC8, 0xA8, 0x4E)
DARK = RGBColor(0x1A, 0x1A, 0x2E)
GRAY = RGBColor(0x66, 0x66, 0x80)
RED = RGBColor(0xCC, 0x33, 0x33)
GREEN = RGBColor(0x1A, 0x87, 0x54)


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = GOLD if level == 1 else DARK
    return h


def add_para(doc, text, bold=False, color=None, size=11, alignment=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color
    if alignment:
        p.alignment = alignment
    return p


def add_key_value(doc, key, value):
    p = doc.add_paragraph()
    k = p.add_run(f"{key}: ")
    k.bold = True
    k.font.size = Pt(11)
    k.font.color.rgb = DARK
    v = p.add_run(str(value))
    v.font.size = Pt(11)
    v.font.color.rgb = GRAY
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(text, style='List Bullet')
    for run in p.runs:
        run.font.size = Pt(10)
    return p


def generate_docx(brief_path, output_path=None):
    """Generate a Word document from a brief JSON file."""
    with open(brief_path) as f:
        data = json.load(f)

    meta = data.get("metadata", {})
    phases = data.get("phases", {})
    company = meta.get("company", "Unknown")
    domain = meta.get("domain", "")
    generated = meta.get("generated_at", "")

    p1 = phases.get("phase_1", {})
    p2 = phases.get("phase_2", {})
    p3 = phases.get("phase_3", {})
    p4 = phases.get("phase_4", {})
    p5 = phases.get("phase_5", {})
    p6 = phases.get("phase_6", {})

    doc = Document()

    # Style setup
    style = doc.styles['Normal']
    style.font.name = 'Arial'
    style.font.size = Pt(11)

    # ── TITLE PAGE ──
    for _ in range(6):
        doc.add_paragraph()

    add_para(doc, "SALES INTELLIGENCE BRIEF", bold=True, size=22, color=DARK, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, company.upper(), bold=True, size=28, color=GOLD, alignment=WD_ALIGN_PARAGRAPH.CENTER)

    subtitle_parts = [domain, p1.get("headquarters", "")]
    subtitle = " | ".join(p for p in subtitle_parts if p)
    if subtitle:
        add_para(doc, subtitle, size=11, color=GRAY, alignment=WD_ALIGN_PARAGRAPH.CENTER)

    for _ in range(4):
        doc.add_paragraph()

    add_para(doc, f"Prepared by Nexus BDR Intelligence System", size=9, color=GRAY, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, f"Generated: {generated[:10] if generated else 'N/A'} | 6-Phase AI Research Pipeline", size=9, color=GRAY, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, f"For: Shareef El-Sissi, Founder & CEO — Terpene Belt Farms", size=10, color=DARK, alignment=WD_ALIGN_PARAGRAPH.CENTER)

    doc.add_page_break()

    # ── EXECUTIVE SUMMARY ──
    if p6 and p6.get("executive_summary"):
        es = p6["executive_summary"]
        add_heading(doc, "Executive Summary")
        if es.get("one_liner"):
            add_para(doc, es["one_liner"], bold=True, size=12)
        add_key_value(doc, "Recommended Brand", es.get("recommended_brand", "TBD"))
        add_key_value(doc, "Priority Score", f"{es.get('priority_score', '?')}/100")
        add_key_value(doc, "Confidence", es.get("confidence", "?"))
        add_key_value(doc, "Expected Annual Value", es.get("annual_value", "?"))
        if es.get("key_insight"):
            add_key_value(doc, "Key Insight", es["key_insight"])
        doc.add_paragraph()

    # ── COMPANY PROFILE ──
    if p1 and len(p1) > 1:
        add_heading(doc, "Company Profile")
        if p1.get("description"):
            add_para(doc, p1["description"])

        add_key_value(doc, "Founded", p1.get("founded", "?"))
        add_key_value(doc, "Headquarters", p1.get("headquarters", "?"))
        add_key_value(doc, "Business Model", p1.get("business_model", "?"))
        add_key_value(doc, "Company Type", p1.get("company_type", "?"))
        add_key_value(doc, "Terpene Relevance", p1.get("terpene_relevance", "?"))

        size = p1.get("size_estimate", {})
        if size:
            add_heading(doc, "Size & Scale", level=2)
            add_key_value(doc, "Employees", size.get("employees", "?"))
            add_key_value(doc, "Revenue Estimate", size.get("revenue_estimate", "?"))
            add_key_value(doc, "Production Scale", size.get("production_scale", "?"))

        products = p1.get("key_products", p1.get("key_products_overview", []))
        if products:
            add_heading(doc, "Key Products", level=2)
            for prod in products:
                if isinstance(prod, str):
                    add_bullet(doc, prod)
                elif isinstance(prod, dict):
                    add_bullet(doc, f"{prod.get('name', '?')} — {prod.get('type', '')} [{prod.get('terpene_relevance', '?')}]")

        news = p1.get("recent_news", [])
        if news:
            add_heading(doc, "Recent News & Signals", level=2)
            for n in news:
                if isinstance(n, dict):
                    add_bullet(doc, f"[{n.get('date', '?')}] {n.get('headline', '?')} — {n.get('significance', '')[:150]}")

        signals = p1.get("growth_signals", [])
        if signals:
            add_heading(doc, "Growth Signals", level=2)
            for s in signals:
                add_bullet(doc, s if isinstance(s, str) else str(s))

    # ── PRODUCTS & TERPENE ANALYSIS ──
    if p2 and len(p2) > 1:
        doc.add_page_break()
        add_heading(doc, "Product & Terpene Analysis")

        terp = p2.get("terpene_analysis", {})
        if terp:
            add_key_value(doc, "Current Terpene Usage", terp.get("current_usage", "?"))
            add_key_value(doc, "Terpene Type", terp.get("terpene_type", "?"))
            add_key_value(doc, "Est. Volume", f"{terp.get('volume_estimate_liters_monthly', '?')} L/month")

            clues = terp.get("supplier_clues", [])
            if clues:
                add_heading(doc, "Supplier Clues", level=2)
                for c in clues:
                    add_bullet(doc, str(c))

        matches = p2.get("product_matches", [])
        if matches:
            add_heading(doc, "Product Match Opportunities", level=2)
            for m in matches:
                if isinstance(m, dict):
                    add_bullet(doc, f"{m.get('their_product', '?')} → {m.get('our_match', '?')}: {m.get('rationale', '')}")

        feedback = p2.get("consumer_feedback", {})
        if feedback and any(feedback.values()):
            add_heading(doc, "Consumer Feedback", level=2)
            if feedback.get("quality_issues"):
                add_key_value(doc, "Quality Issues", feedback["quality_issues"])
            if feedback.get("opportunities"):
                add_key_value(doc, "Opportunities", feedback["opportunities"])

    # ── PEOPLE ──
    if p3 and len(p3) > 1:
        dms = p3.get("decision_makers", [])
        if dms:
            add_heading(doc, "Decision Makers", level=2)
            for dm in dms:
                if isinstance(dm, dict):
                    add_para(doc, f"{dm.get('name', '?')} — {dm.get('title', '?')}", bold=True)
                    add_key_value(doc, "Role in Purchase", dm.get("role_in_purchase", "?"))
                    hooks = dm.get("personalization_hooks", [])
                    if hooks:
                        for h in hooks[:3]:
                            add_bullet(doc, str(h) if isinstance(h, str) else h.get("hook", str(h)))

    # ── COMPETITIVE INTELLIGENCE ──
    if p4 and len(p4) > 1:
        doc.add_page_break()
        add_heading(doc, "Competitive & Supplier Intelligence")

        supplier = p4.get("current_supplier", p4.get("current_supplier_assessment", {}))
        if supplier:
            add_key_value(doc, "Most Likely Supplier", supplier.get("most_likely", supplier.get("most_likely_supplier", "?")))
            evidence = supplier.get("evidence", [])
            if evidence:
                add_heading(doc, "Evidence", level=2)
                for e in evidence:
                    add_bullet(doc, str(e))

            barriers = supplier.get("switching_barriers", [])
            if barriers:
                add_heading(doc, "Switching Barriers", level=2)
                for b in barriers:
                    add_bullet(doc, str(b))

            triggers = supplier.get("switching_triggers", [])
            if triggers:
                add_heading(doc, "Switching Triggers", level=2)
                for t in triggers:
                    add_bullet(doc, str(t))

        displacement = p4.get("displacement", p4.get("displacement_strategy", {}))
        if displacement:
            add_heading(doc, "Displacement Strategy", level=2)
            add_key_value(doc, "Primary Angle", displacement.get("primary_angle", "?"))
            angles = displacement.get("supporting_angles", [])
            for a in angles:
                add_bullet(doc, str(a))

        pricing = p4.get("pricing", p4.get("pricing_intelligence", {}))
        if pricing:
            add_heading(doc, "Pricing Analysis", level=2)
            add_key_value(doc, "Their Current Cost", pricing.get("their_likely_cost_per_liter", pricing.get("their_likely_current_cost", "?")))
            add_key_value(doc, "Our Price", pricing.get("our_price", pricing.get("our_pricing_tier", "?")))
            add_key_value(doc, "Savings", pricing.get("savings", pricing.get("advantage_or_gap", "?")))

    # ── FINANCIAL MODEL ──
    if p5 and len(p5) > 1:
        add_heading(doc, "Financial Model")
        deal = p5.get("deal_model", {})
        for scenario in ["conservative", "likely", "upside"]:
            s = deal.get(scenario, {})
            if s:
                add_para(doc, f"{scenario.upper()}: {s.get('desc', '?')}", bold=True)
                add_key_value(doc, "Annual Value", s.get("annual_value", "?"))

        if deal.get("weighted_annual"):
            add_para(doc, f"Weighted Annual: {deal['weighted_annual']}", bold=True, size=13, color=GOLD)

    # ── OUTREACH SEQUENCE ──
    if p6 and p6.get("outreach_sequence"):
        doc.add_page_break()
        add_heading(doc, "Outreach Sequence")

        for touch in p6["outreach_sequence"]:
            if isinstance(touch, dict):
                add_heading(doc, f"Touch {touch.get('touch', '?')}: {touch.get('channel', '?')} (Day {touch.get('day', '?')})", level=2)
                add_key_value(doc, "Target", touch.get("target", "?"))
                if touch.get("subject"):
                    add_key_value(doc, "Subject", touch["subject"])
                if touch.get("body"):
                    add_para(doc, touch["body"])
                if touch.get("message"):
                    add_para(doc, touch["message"])

    # ── ACTION PLAN ──
    if p6 and p6.get("action_plan"):
        add_heading(doc, "Action Plan")
        plan = p6["action_plan"]
        for timeframe in ["today", "this_week", "next_30_days"]:
            items = plan.get(timeframe, [])
            if items:
                add_heading(doc, timeframe.replace("_", " ").title(), level=2)
                for item in items:
                    add_bullet(doc, str(item))

    # ── SAVE ──
    if not output_path:
        slug = re.sub(r'[^a-z0-9]', '_', company[:30].lower())
        output_path = OUTPUT_DIR / f"brief_{slug}.docx"

    doc.save(str(output_path))
    print(f"  📄 Generated: {output_path}")
    return output_path


def main():
    parser = argparse.ArgumentParser(description="Brief → Word Document Generator")
    parser.add_argument("--brief", help="Path to brief JSON file (supports glob)")
    parser.add_argument("--output", help="Custom output path")
    parser.add_argument("--all", action="store_true", help="Convert all briefs")
    args = parser.parse_args()

    if args.all:
        briefs = sorted(BRIEFS_DIR.glob("brief_*.json"))
        print(f"  Found {len(briefs)} briefs")
        for b in briefs:
            generate_docx(b)
    elif args.brief:
        # Support glob
        matches = glob.glob(args.brief)
        if not matches:
            print(f"  ❌ No file found: {args.brief}")
            return
        for m in matches:
            generate_docx(m, args.output)
    else:
        parser.print_help()


if __name__ == "__main__":
    main()
